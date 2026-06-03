import os
import sqlite3
from uuid import uuid4

from flask import Blueprint, render_template, request, redirect, url_for, flash
from flask_login import login_required
from werkzeug.utils import secure_filename

from database import get_db_connection

questions_bp = Blueprint('questions', __name__)

OPTION_KEYS = ('a', 'b', 'c', 'd', 'e')
DIFFICULTY_CHOICES = ('easy', 'hard')
DIFFICULTY_LABELS = {'easy': 'Ýönekeý', 'hard': 'Çylşyrymly'}


def _project_root() -> str:
    return os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))


def _ensure_dir(path: str) -> None:
    os.makedirs(path, exist_ok=True)


def _delete_image(rel_path: str | None) -> None:
    if not rel_path:
        return
    abs_path = os.path.join(_project_root(), 'static', rel_path.replace('/', os.sep))
    if os.path.isfile(abs_path):
        try:
            os.remove(abs_path)
        except OSError:
            pass


def _save_image(file_storage, rel_dir: str) -> str | None:
    if not file_storage or not getattr(file_storage, 'filename', None):
        return None
    if file_storage.filename == '':
        return None

    static_dir = os.path.join(_project_root(), 'static')
    abs_dir = os.path.join(static_dir, rel_dir)
    _ensure_dir(abs_dir)

    original_ext = os.path.splitext(file_storage.filename)[1].lower()
    if not original_ext:
        original_ext = '.img'

    filename = secure_filename(f'{uuid4().hex}{original_ext}')
    abs_path = os.path.join(abs_dir, filename)
    file_storage.save(abs_path)

    return os.path.join(rel_dir, filename).replace('\\', '/')


def _load_question(db, question_id: int):
    question = db.execute('SELECT * FROM questions WHERE id = ?', (question_id,)).fetchone()
    if not question:
        return None, None

    options_rows = db.execute(
        'SELECT option_key, option_text, option_image FROM question_options WHERE question_id = ?',
        (question_id,),
    ).fetchall()
    options = {row['option_key']: dict(row) for row in options_rows}

    for k in OPTION_KEYS:
        options.setdefault(k, {'option_key': k, 'option_text': None, 'option_image': None})

    return question, options


def _load_subjects_for_dropdown(db, current_subject_id=None):
    active_subjects = db.execute(
        'SELECT id, name, is_active, answer_count FROM subjects WHERE is_active = 1 ORDER BY name ASC'
    ).fetchall()
    if current_subject_id is not None:
        current = db.execute(
            'SELECT id, name, is_active, answer_count FROM subjects WHERE id = ?',
            (current_subject_id,),
        ).fetchone()
        if current and not any(s['id'] == current['id'] for s in active_subjects):
            active_subjects = list(active_subjects) + [current]
            active_subjects.sort(key=lambda s: s['name'])

    if not active_subjects:
        active_subjects = db.execute(
            'SELECT id, name, is_active, answer_count FROM subjects ORDER BY name ASC'
        ).fetchall()

    return active_subjects


def _subject_answer_counts(subjects) -> dict:
    return {str(s['id']): s['answer_count'] for s in subjects}


def _get_answer_count(db, subject_id: int) -> int:
    row = db.execute('SELECT answer_count FROM subjects WHERE id = ?', (subject_id,)).fetchone()
    return row['answer_count'] if row else 5


def _render_form(question, options, subjects, draft, current_subject_id):
    return render_template(
        'questions/form.html',
        question=question,
        options=options,
        subjects=subjects,
        draft=draft,
        current_subject_id=current_subject_id,
        option_keys=OPTION_KEYS,
        difficulty_choices=DIFFICULTY_CHOICES,
        difficulty_labels=DIFFICULTY_LABELS,
        subject_answer_counts=_subject_answer_counts(subjects),
    )


_DIFFICULTY_MAP = {
    'easy': 'easy', 'ýönekeý': 'easy', 'йонекей': 'easy', 'простой': 'easy', '1': 'easy',
    'hard': 'hard', 'çylşyrymly': 'hard', 'чылшырымлы': 'hard', 'сложный': 'hard', '2': 'hard',
}


def _parse_questions_excel(file_bytes: bytes):
    from io import BytesIO
    import openpyxl

    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as e:
        return None, f'Faýly açmak başartmady: {e}'

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(rows) < 2:
        return None, 'Faýlda maglumat ýok'

    headers = [str(c).strip().lower() if c is not None else '' for c in rows[0]]

    def _col(names):
        for name in names:
            for i, h in enumerate(headers):
                if name in h:
                    return i
        return None

    def _col_exact(names):
        for name in names:
            for i, h in enumerate(headers):
                if h == name:
                    return i
        return None

    col_subject  = _col(['ders', 'subject', 'предмет'])
    col_text     = _col(['sorag', 'question', 'вопрос', 'текст', 'text'])
    col_diff     = _col(['kynç', 'diff', 'тип', 'тяж', 'type', 'сложн'])
    col_correct  = _col(['dogry', 'correct', 'jogap', 'ответ', 'answer'])
    col_opts     = {k: _col_exact([k, f'{k})']) for k in OPTION_KEYS}

    missing = []
    if col_subject is None:  missing.append('Ders')
    if col_text is None:     missing.append('Sorag')
    if col_diff is None:     missing.append('Kynçylyk')
    if col_correct is None:  missing.append('Dogry jogap')
    if col_opts['a'] is None: missing.append('A')
    if col_opts['b'] is None: missing.append('B')
    if missing:
        return None, f'Sütünler tapylmady: {", ".join(missing)}. Birinji setirde at barmy barlaň.'

    def _cell(row, idx):
        if idx is None or idx >= len(row):
            return ''
        v = row[idx]
        return str(v).strip() if v is not None else ''

    records = []
    for row_num, row in enumerate(rows[1:], 2):
        if not any(c for c in row if c is not None and str(c).strip()):
            continue

        subject_name = _cell(row, col_subject)
        question_text = _cell(row, col_text) or None
        diff_raw = _cell(row, col_diff).lower()
        correct_raw = _cell(row, col_correct)

        if not subject_name:
            records.append({'ok': False, 'row_num': row_num, 'error': 'ders ady boş'})
            continue

        difficulty = _difficulty_map_lookup(diff_raw)
        if difficulty is None:
            records.append({'ok': False, 'row_num': row_num, 'error': f'nämälim kynçylyk «{diff_raw}»'})
            continue

        correct_keys = [k.strip().lower() for k in correct_raw.replace(',', ' ').split() if k.strip()]
        correct_keys = [k for k in correct_keys if k in OPTION_KEYS]
        if not correct_keys:
            records.append({'ok': False, 'row_num': row_num, 'error': f'dogry jogap ýok ýa-da nädogry «{correct_raw}»'})
            continue

        options = {}
        for k in OPTION_KEYS:
            txt = _cell(row, col_opts[k]) if col_opts[k] is not None else ''
            options[k] = txt or None

        records.append({
            'ok': True,
            'row_num': row_num,
            'subject_name': subject_name,
            'question_text': question_text,
            'difficulty': difficulty,
            'correct_keys': correct_keys,
            'options': options,
        })

    return records, None


def _difficulty_map_lookup(raw: str) -> str | None:
    for key, val in _DIFFICULTY_MAP.items():
        if key in raw:
            return val
    return None


@questions_bp.route('/import', methods=['POST'])
@login_required
def import_excel():
    file = request.files.get('file')
    if not file or file.filename == '':
        flash('Faýl saýlaň', 'danger')
        return redirect(url_for('questions.index'))

    if not file.filename.lower().endswith(('.xlsx', '.xls')):
        flash('Diňe .xlsx we .xls faýllar goldanylýar', 'danger')
        return redirect(url_for('questions.index'))

    records, error = _parse_questions_excel(file.read())
    if error:
        flash(f'Import ýalňyşlygy: {error}', 'danger')
        return redirect(url_for('questions.index'))

    db = get_db_connection()
    subjects_cache = {}
    imported = 0
    errors = []

    for rec in records:
        if not rec['ok']:
            errors.append(f"Setir {rec['row_num']}: {rec['error']}")
            continue

        sname = rec['subject_name']
        if sname not in subjects_cache:
            row = db.execute(
                'SELECT id, answer_count FROM subjects WHERE lower(name) = lower(?)', (sname,)
            ).fetchone()
            subjects_cache[sname] = dict(row) if row else None

        n = rec['row_num']

        subject = subjects_cache[sname]
        if subject is None:
            errors.append(f'Setir {n}: ders tapylmady «{sname}»')
            continue

        answer_count = subject['answer_count']
        active_keys = OPTION_KEYS[:answer_count]

        valid_correct = [k for k in rec['correct_keys'] if k in active_keys]
        if not valid_correct:
            errors.append(
                f'Setir {n}: «{sname}» üçin dogry jogap {[k.upper() for k in active_keys]} '
                f'içinde bolmaly, alnan: {[k.upper() for k in rec["correct_keys"]]}'
            )
            continue

        missing_opts = [k for k in active_keys if not rec['options'].get(k)]
        if missing_opts:
            errors.append(
                f'Setir {n}: {[k.upper() for k in missing_opts]} wariantlary boş '
                f'(«{sname}» üçin {answer_count} wariant gerek)'
            )
            continue

        if rec['question_text']:
            exists = db.execute(
                'SELECT id FROM questions WHERE subject_id = ? AND question_text = ?',
                (subject['id'], rec['question_text']),
            ).fetchone()
            if exists:
                errors.append(f"Setir {n}: bu sorag eýýäm bar (ID {exists['id']}), geçildi")
                continue

        correct_str = ','.join(sorted(valid_correct))
        try:
            cursor = db.cursor()
            cursor.execute(
                'INSERT INTO questions (subject_id, question_text, correct_option, difficulty) VALUES (?, ?, ?, ?)',
                (subject['id'], rec['question_text'], correct_str, rec['difficulty']),
            )
            qid = cursor.lastrowid
            for k in active_keys:
                cursor.execute(
                    'INSERT INTO question_options (question_id, option_key, option_text) VALUES (?, ?, ?)',
                    (qid, k, rec['options'].get(k)),
                )
            db.commit()
            imported += 1
        except Exception as e:
            db.rollback()
            errors.append(f'Setir {n}: DB ýalňyşlygy: {e}')

    db.close()

    if imported:
        flash(f'{imported} sorag import edildi', 'success')
    for e in errors:
        flash(e, 'warning')

    return redirect(url_for('questions.index'))


@questions_bp.route('/')
@login_required
def index():
    db = get_db_connection()
    rows = db.execute(
        '''
        SELECT
            q.id, q.is_active, q.subject_id,
            s.name AS subject_name,
            q.question_text, q.question_image,
            q.correct_option, q.difficulty
        FROM questions q
        JOIN subjects s ON q.subject_id = s.id
        ORDER BY s.name ASC, q.id DESC
        '''
    ).fetchall()
    db.close()
    return render_template('questions/list.html', questions=rows, difficulty_labels=DIFFICULTY_LABELS)


@questions_bp.route('/create', methods=['GET', 'POST'])
@login_required
def create():
    if request.method == 'POST':
        db = get_db_connection()

        subject_id_raw = request.form.get('subject_id')
        try:
            subject_id = int(subject_id_raw)
        except (TypeError, ValueError):
            subject_id = None

        subject = None
        if subject_id is not None:
            subject = db.execute('SELECT id, is_active, answer_count FROM subjects WHERE id = ?', (subject_id,)).fetchone()

        if not subject or subject['is_active'] != 1:
            subjects = _load_subjects_for_dropdown(db)
            db.close()
            flash('Saýlanan ders tapylmady ýa-da işjeň däl', 'danger')
            return _render_form(
                None,
                {k: {'option_key': k, 'option_text': None, 'option_image': None} for k in OPTION_KEYS},
                subjects,
                {'question_text': request.form.get('question_text', ''), 'correct_options_list': request.form.getlist('correct_options'), 'difficulty': request.form.get('difficulty', 'easy')},
                subject_id,
            )

        answer_count = subject['answer_count']
        active_keys = OPTION_KEYS[:answer_count]

        question_text = request.form.get('question_text', '').strip() or None
        difficulty = request.form.get('difficulty', 'easy')
        if difficulty not in DIFFICULTY_CHOICES:
            difficulty = 'easy'

        correct_options_raw = [k for k in request.form.getlist('correct_options') if k in active_keys]

        def _make_draft():
            return {
                'question_text': request.form.get('question_text', ''),
                'correct_options_list': correct_options_raw,
                'difficulty': difficulty,
            }

        def _make_options_draft():
            return {
                k: {'option_key': k, 'option_text': request.form.get(f'option_{k}_text', '').strip() or None, 'option_image': None}
                for k in OPTION_KEYS
            }

        if not correct_options_raw:
            subjects = _load_subjects_for_dropdown(db, current_subject_id=subject_id)
            db.close()
            flash('Iň bolmanda bir dogry görnüş seçmeli', 'danger')
            return _render_form(None, _make_options_draft(), subjects, _make_draft(), subject_id)

        question_image_file = request.files.get('question_image')
        has_new_question_image = bool(question_image_file and question_image_file.filename)

        if not question_text and not has_new_question_image:
            subjects = _load_subjects_for_dropdown(db, current_subject_id=subject_id)
            db.close()
            flash('Soragyň teksti ýa-da surady bolmaly', 'danger')
            return _render_form(None, _make_options_draft(), subjects, _make_draft(), subject_id)

        option_final_text = {}
        option_new_files = {}
        option_has_image_new = {}
        errors = []

        for k in active_keys:
            opt_text = request.form.get(f'option_{k}_text', '').strip() or None
            opt_file = request.files.get(f'option_{k}_image')
            has_new_image = bool(opt_file and opt_file.filename)

            if not opt_text and not has_new_image:
                errors.append(f'Görnüş {k.upper()} surady ýa-da teksti bolmaly')

            option_final_text[k] = opt_text
            option_new_files[k] = opt_file
            option_has_image_new[k] = has_new_image

        if errors:
            subjects = _load_subjects_for_dropdown(db, current_subject_id=subject_id)
            db.close()
            for e in errors:
                flash(e, 'danger')
            return _render_form(None, _make_options_draft(), subjects, _make_draft(), subject_id)

        question_image_path = _save_image(question_image_file, os.path.join('uploads', 'questions'))
        option_image_paths = {k: _save_image(option_new_files[k], os.path.join('uploads', 'questions')) for k in active_keys}

        correct_option_str = ','.join(sorted(correct_options_raw))

        try:
            cursor = db.cursor()
            cursor.execute(
                'INSERT INTO questions (subject_id, question_text, question_image, correct_option, difficulty) VALUES (?, ?, ?, ?, ?)',
                (subject_id, question_text, question_image_path, correct_option_str, difficulty),
            )
            question_id = cursor.lastrowid

            for k in active_keys:
                cursor.execute(
                    'INSERT INTO question_options (question_id, option_key, option_text, option_image) VALUES (?, ?, ?, ?)',
                    (question_id, k, option_final_text[k], option_image_paths[k]),
                )
            db.commit()
            db.close()
            flash('Sorag üstünikli goşulan', 'success')
            return redirect(url_for('questions.index'))
        except sqlite3.IntegrityError:
            db.rollback()
            db.close()
            flash('Ýatda saklamada ýalňyş çykdy: meýdançalaryň bahalaryny barlaň', 'danger')

    db = get_db_connection()
    subjects = _load_subjects_for_dropdown(db)
    db.close()
    options = {k: {'option_key': k, 'option_text': None, 'option_image': None} for k in OPTION_KEYS}
    return _render_form(
        None,
        options,
        subjects,
        {'question_text': '', 'correct_options_list': [], 'difficulty': 'easy'},
        None,
    )


@questions_bp.route('/edit/<int:id>', methods=['GET', 'POST'])
@login_required
def edit(id: int):
    db = get_db_connection()
    question, options = _load_question(db, id)
    if not question:
        db.close()
        flash('Sorag tapylmady', 'danger')
        return redirect(url_for('questions.index'))

    if request.method == 'POST':
        subject_id_raw = request.form.get('subject_id')
        try:
            subject_id = int(subject_id_raw)
        except (TypeError, ValueError):
            subject_id = None

        subject = None
        if subject_id is not None:
            subject = db.execute('SELECT id, answer_count FROM subjects WHERE id = ?', (subject_id,)).fetchone()

        if not subject:
            subjects = _load_subjects_for_dropdown(db, current_subject_id=question['subject_id'])
            db.close()
            flash('Saýlanan ders tapylmady', 'danger')
            options_draft = {k: {'option_key': k, 'option_text': request.form.get(f'option_{k}_text', '').strip() or None, 'option_image': options[k].get('option_image')} for k in OPTION_KEYS}
            return _render_form(question, options_draft, subjects, {'question_text': request.form.get('question_text', ''), 'correct_options_list': request.form.getlist('correct_options'), 'difficulty': request.form.get('difficulty', 'easy')}, question['subject_id'])

        answer_count = subject['answer_count']
        active_keys = OPTION_KEYS[:answer_count]

        question_text = request.form.get('question_text', '').strip() or None
        difficulty = request.form.get('difficulty', 'easy')
        if difficulty not in DIFFICULTY_CHOICES:
            difficulty = 'easy'

        correct_options_raw = [k for k in request.form.getlist('correct_options') if k in active_keys]

        def _make_draft():
            return {
                'question_text': request.form.get('question_text', ''),
                'correct_options_list': correct_options_raw,
                'difficulty': difficulty,
            }

        def _make_options_draft():
            return {
                k: {'option_key': k, 'option_text': request.form.get(f'option_{k}_text', '').strip() or None, 'option_image': options[k].get('option_image')}
                for k in OPTION_KEYS
            }

        if not correct_options_raw:
            subjects = _load_subjects_for_dropdown(db, current_subject_id=question['subject_id'])
            db.close()
            flash('Iň bolmanda bir dogry görnüş seçmeli', 'danger')
            return _render_form(question, _make_options_draft(), subjects, _make_draft(), question['subject_id'])

        question_image_file = request.files.get('question_image')
        has_new_question_image = bool(question_image_file and question_image_file.filename)
        final_question_image_present = has_new_question_image or bool(question['question_image'])

        if not question_text and not final_question_image_present:
            subjects = _load_subjects_for_dropdown(db, current_subject_id=question['subject_id'])
            db.close()
            flash('Soragyň teksti ýa-da surady bolmaly', 'danger')
            return _render_form(question, _make_options_draft(), subjects, _make_draft(), question['subject_id'])

        new_option_files = {}
        final_option_text = {}
        final_option_image_paths = {}
        errors = []

        for k in active_keys:
            opt_text = request.form.get(f'option_{k}_text', '').strip() or None
            opt_file = request.files.get(f'option_{k}_image')
            has_new_image = bool(opt_file and opt_file.filename)
            existing_image = options[k].get('option_image')

            if not opt_text and not has_new_image and not existing_image:
                errors.append(f'Görnüş {k.upper()} surady ýa-da teksti bolmalydyr')

            final_option_text[k] = opt_text
            new_option_files[k] = opt_file
            final_option_image_paths[k] = existing_image

        if errors:
            subjects = _load_subjects_for_dropdown(db, current_subject_id=question['subject_id'])
            db.close()
            for e in errors:
                flash(e, 'danger')
            return _render_form(question, _make_options_draft(), subjects, _make_draft(), question['subject_id'])

        new_question_image_path = None
        if has_new_question_image:
            _delete_image(question['question_image'])
            new_question_image_path = _save_image(question_image_file, os.path.join('uploads', 'questions'))

        for k in active_keys:
            if new_option_files[k] and new_option_files[k].filename:
                _delete_image(options[k].get('option_image'))
                final_option_image_paths[k] = _save_image(new_option_files[k], os.path.join('uploads', 'questions'))

        correct_option_str = ','.join(sorted(correct_options_raw))

        try:
            db.execute(
                'UPDATE questions SET subject_id=?, question_text=?, question_image=?, correct_option=?, difficulty=? WHERE id=?',
                (subject_id, question_text, new_question_image_path or question['question_image'], correct_option_str, difficulty, id),
            )

            placeholders = ','.join('?' for _ in active_keys)
            db.execute(
                f'DELETE FROM question_options WHERE question_id=? AND option_key NOT IN ({placeholders})',
                [id] + list(active_keys),
            )

            for k in active_keys:
                db.execute(
                    '''
                    INSERT INTO question_options (question_id, option_key, option_text, option_image)
                    VALUES (?, ?, ?, ?)
                    ON CONFLICT(question_id, option_key)
                    DO UPDATE SET option_text=excluded.option_text, option_image=excluded.option_image
                    ''',
                    (id, k, final_option_text[k], final_option_image_paths[k]),
                )

            db.commit()
            db.close()
            flash('Sorag üstünikli üýtgedildi', 'success')
            return redirect(url_for('questions.index'))
        except sqlite3.IntegrityError:
            db.rollback()
            db.close()
            flash('Ýatda saklamada ýalňyş çykdy: meýdançalaryň bahalaryny barlaň', 'danger')

    subjects = _load_subjects_for_dropdown(db, current_subject_id=question['subject_id'])
    db.close()
    correct_options_list = question['correct_option'].split(',')
    return _render_form(
        question,
        options,
        subjects,
        {'question_text': '', 'correct_options_list': correct_options_list, 'difficulty': question['difficulty']},
        question['subject_id'],
    )


@questions_bp.route('/delete/<int:id>', methods=['POST'])
@login_required
def delete(id: int):
    db = get_db_connection()
    question = db.execute('SELECT question_image FROM questions WHERE id = ?', (id,)).fetchone()
    if not question:
        db.close()
        flash('Sorag tapylmady', 'danger')
        return redirect(url_for('questions.index'))

    used = db.execute(
        'SELECT COUNT(*) FROM test_instance_questions WHERE question_id = ?', (id,)
    ).fetchone()[0]
    if used:
        db.close()
        flash(f'Soragy ýok etmek mümkin däl — ol {used} test nusgasynda ulanylýar', 'danger')
        return redirect(url_for('questions.index'))

    opt_images = db.execute(
        'SELECT option_image FROM question_options WHERE question_id = ?', (id,)
    ).fetchall()
    db.execute('DELETE FROM questions WHERE id = ?', (id,))
    db.commit()
    _delete_image(question['question_image'])
    for opt in opt_images:
        _delete_image(opt['option_image'])
    db.close()
    flash('Sorag ýok edildi', 'info')
    return redirect(url_for('questions.index'))


@questions_bp.route('/toggle/<int:id>', methods=['POST'])
@login_required
def toggle(id: int):
    db = get_db_connection()
    db.execute('UPDATE questions SET is_active = 1 - is_active WHERE id = ?', (id,))
    db.commit()
    db.close()
    flash('Soragyň ýagdaýy üýtgedildi', 'info')
    return redirect(url_for('questions.index'))
